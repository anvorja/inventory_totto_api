"""Lectura de los Excel que maneja la tienda.

Se reconocen dos tipos de archivo:

* **Existencias** (export del sistema, se actualiza ~2 veces al día): columnas
  `Punto de Venta`, `Código Producto Largo`, `Disponible` y opcionalmente `UND`,
  `Nombre Producto Largo`, `Talla`, `CodColor`, `DsColor`.
* **Maestro de códigos de barras**: relaciona la referencia Totto con el EAN impreso en
  la etiqueta. Convención: `BARCODE` = EAN-13 y `Código Producto Largo` (o `SKU`) =
  referencia. Si un archivo trae los encabezados invertidos, se detecta por el contenido,
  se corrige y se informa (`columns_swapped`).

El parser es deliberadamente tolerante: busca la fila de encabezados en las primeras
filas, ignora mayúsculas/tildes/espacios, descarta la fila "Total" y notas al pie.
"""

from __future__ import annotations

import io
import re
import unicodedata
import warnings
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from enum import StrEnum
from typing import Any

from openpyxl import load_workbook

HEADER_SEARCH_ROWS = 20
EAN_RE = re.compile(r"^\d{8,14}$")


class ImportKind(StrEnum):
    STOCK = "stock"
    CATALOG = "catalog"


class XlsxFormatError(ValueError):
    """El archivo no tiene la forma esperada. El mensaje se muestra al usuario."""


# ----------------------------------------------------------------------------- utilidades


def normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]+", " ", text.upper()).strip()


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def clean_reference(value: Any) -> str | None:
    text = clean_text(value)
    return text.upper().replace(" ", "") if text else None


def clean_ean(value: Any) -> str | None:
    """Normaliza un EAN venga como texto, entero o flotante (7.70468E+12)."""
    if value is None:
        return None
    if isinstance(value, float):
        if not value.is_integer():
            return None
        value = int(value)
    text = re.sub(r"\s+", "", str(value))
    if text.endswith(".0"):
        text = text[:-2]
    return text if EAN_RE.match(text) else None


def to_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return round(value)
    text = str(value).strip().replace(" ", "")
    # "1.234" o "1,234" como separador de miles; "3,0" o "3.0" como decimal.
    if re.fullmatch(r"-?\d{1,3}([.,]\d{3})+", text):
        text = re.sub(r"[.,]", "", text)
    text = text.replace(",", ".")
    try:
        return round(float(text))
    except ValueError:
        return None


def split_store(value: Any) -> tuple[str, str] | None:
    """'FQ95 - TOTTO VILLA GORGONA' -> ('FQ95', 'TOTTO VILLA GORGONA')."""
    text = clean_text(value)
    if not text:
        return None
    code, sep, name = text.partition(" - ")
    code = code.strip().upper()
    if not sep:
        return code[:20], text
    return code[:20], name.strip() or text


# ------------------------------------------------------------------------- lectura cruda


def read_rows(content: bytes) -> list[tuple[Any, ...]]:
    try:
        with warnings.catch_warnings():
            # Los exports del sistema no traen estilos por defecto; es inofensivo.
            warnings.simplefilter("ignore", UserWarning)
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza varios tipos según el daño del archivo
        raise XlsxFormatError(
            "No se pudo abrir el archivo. Asegúrate de que sea un Excel .xlsx válido."
        ) from exc
    try:
        ws = wb.worksheets[0]
        return [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


@dataclass
class HeaderMatch:
    row_index: int
    columns: dict[str, int]


def find_header(
    rows: Sequence[Sequence[Any]], aliases: dict[str, Iterable[str]], required: Iterable[str]
) -> HeaderMatch | None:
    required = set(required)
    normalized_aliases = {
        key: {normalize_header(a) for a in names} for key, names in aliases.items()
    }
    for i, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        headers = [normalize_header(c) for c in row]
        columns: dict[str, int] = {}
        for key, names in normalized_aliases.items():
            for col, header in enumerate(headers):
                if header in names:
                    columns[key] = col
                    break
        if required <= columns.keys():
            return HeaderMatch(i, columns)
    return None


def cell(row: Sequence[Any], col: int | None) -> Any:
    if col is None or col >= len(row):
        return None
    return row[col]


# ------------------------------------------------------------------------- existencias

STOCK_ALIASES = {
    "store": ["Punto de Venta", "Tienda", "Almacen"],
    "reference": ["Código Producto Largo", "Codigo Producto", "Referencia"],
    "business_unit": ["UND", "Unidad de Negocio"],
    "name": ["Nombre Producto Largo", "Nombre Producto", "Descripcion"],
    "size": ["Talla"],
    "color_code": ["CodColor", "Cod Color"],
    "color_name": ["DsColor", "Color", "Descripcion Color"],
    "quantity": ["Disponible", "Existencia", "Existencias", "Cantidad", "Stock"],
}
STOCK_REQUIRED = ("store", "reference", "quantity")


@dataclass
class StockLine:
    reference: str
    quantity: int
    name: str | None = None
    business_unit: str | None = None
    size: str | None = None
    color_code: str | None = None
    color_name: str | None = None


@dataclass
class StoreStock:
    code: str
    name: str
    lines: dict[str, StockLine] = field(default_factory=dict)

    @property
    def total_units(self) -> int:
        return sum(line.quantity for line in self.lines.values())


@dataclass
class StockParseResult:
    stores: list[StoreStock]
    skipped_rows: int


def parse_stock(rows: Sequence[Sequence[Any]]) -> StockParseResult:
    header = find_header(rows, STOCK_ALIASES, STOCK_REQUIRED)
    if header is None:
        raise XlsxFormatError(
            "No encontré las columnas 'Punto de Venta', 'Código Producto Largo' y 'Disponible'."
        )
    c = header.columns
    stores: dict[str, StoreStock] = {}
    skipped = 0
    for row in rows[header.row_index + 1 :]:
        store = split_store(cell(row, c["store"]))
        reference = clean_reference(cell(row, c["reference"]))
        quantity = to_int(cell(row, c["quantity"]))
        if not store or not reference or quantity is None:
            if any(v not in (None, "") for v in row):
                skipped += 1  # fila "Total", notas de filtros, filas incompletas
            continue
        bucket = stores.setdefault(store[0], StoreStock(code=store[0], name=store[1]))
        existing = bucket.lines.get(reference)
        if existing:
            existing.quantity += quantity
            continue
        bucket.lines[reference] = StockLine(
            reference=reference,
            quantity=quantity,
            name=clean_text(cell(row, c.get("name"))),
            business_unit=clean_text(cell(row, c.get("business_unit"))),
            size=clean_text(cell(row, c.get("size"))),
            color_code=clean_text(cell(row, c.get("color_code"))),
            color_name=clean_text(cell(row, c.get("color_name"))),
        )
    if not stores:
        raise XlsxFormatError("El archivo de existencias no tiene filas con productos.")
    return StockParseResult(stores=list(stores.values()), skipped_rows=skipped)


# --------------------------------------------------------------- maestro de códigos

# Convención: BARCODE es el EAN-13 impreso en la etiqueta (lo que se escanea) y
# "Código Producto Largo" (o SKU) es la referencia Totto.
CATALOG_ALIASES = {
    "ean": ["BARCODE", "EAN", "EAN13", "EAN 13", "Código de Barras", "Codigo Barras"],
    "reference": ["Código Producto Largo", "Referencia", "SKU", "Codigo"],
    "name": ["NOMBRE PRODUCTO", "Nombre Producto Largo", "Nombre", "Descripcion"],
}
CATALOG_REQUIRED = ("ean", "reference")


@dataclass
class CatalogLine:
    reference: str | None
    ean: str
    name: str | None


@dataclass
class CatalogParseResult:
    lines: list[CatalogLine]
    skipped_rows: int
    # True si los encabezados venían invertidos (p. ej. BARCODE con referencias y SKU con
    # EAN, como el BarrcodeFQ95.xlsx original) y se corrigió leyendo el contenido.
    columns_swapped: bool = False


def _ean_ratio(rows: Sequence[Sequence[Any]], col: int) -> float:
    values = [cell(r, col) for r in rows[:200] if cell(r, col) not in (None, "")]
    if not values:
        return 0.0
    return sum(1 for v in values if clean_ean(v)) / len(values)


def parse_catalog(rows: Sequence[Sequence[Any]]) -> CatalogParseResult:
    header = find_header(rows, CATALOG_ALIASES, CATALOG_REQUIRED)
    if header is None:
        raise XlsxFormatError(
            "No encontré las columnas 'BARCODE' (EAN) y 'Código Producto Largo' del maestro "
            "de códigos."
        )
    body = rows[header.row_index + 1 :]
    ean_col, ref_col = header.columns["ean"], header.columns["reference"]
    # Se confía en el contenido más que en el encabezado: el EAN es la columna de dígitos.
    swapped = _ean_ratio(body, ref_col) > _ean_ratio(body, ean_col)
    if swapped:
        ean_col, ref_col = ref_col, ean_col
    if _ean_ratio(body, ean_col) < 0.5:
        raise XlsxFormatError("No encontré una columna con códigos de barras EAN válidos.")
    name_col = header.columns.get("name")

    lines: dict[str, CatalogLine] = {}
    skipped = 0
    for row in body:
        ean = clean_ean(cell(row, ean_col))
        if not ean:
            if any(v not in (None, "") for v in row):
                skipped += 1
            continue
        lines[ean] = CatalogLine(
            reference=clean_reference(cell(row, ref_col)),
            ean=ean,
            name=clean_text(cell(row, name_col)),
        )
    if not lines:
        raise XlsxFormatError("El maestro de códigos no tiene filas válidas.")
    return CatalogParseResult(
        lines=list(lines.values()), skipped_rows=skipped, columns_swapped=swapped
    )


# ------------------------------------------------------------------------ detección


def detect_kind(rows: Sequence[Sequence[Any]]) -> ImportKind:
    if find_header(rows, STOCK_ALIASES, STOCK_REQUIRED):
        return ImportKind.STOCK
    if find_header(rows, CATALOG_ALIASES, CATALOG_REQUIRED):
        return ImportKind.CATALOG
    raise XlsxFormatError(
        "No reconozco este Excel. Sube el reporte de existencias (Punto de Venta, Código "
        "Producto Largo, Disponible) o el maestro de códigos (BARCODE, Código Producto Largo, "
        "NOMBRE PRODUCTO)."
    )
