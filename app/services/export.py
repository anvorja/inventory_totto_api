"""Genera el Excel de resultados de un conteo."""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import CountSession, StockSnapshot
from app.services.comparison import ComparisonLine, ComparisonSummary, LineStatus

STATUS_LABEL = {
    LineStatus.OK: "Cuadrado",
    LineStatus.MISSING: "Faltante",
    LineStatus.SURPLUS: "Sobrante",
    LineStatus.UNEXPECTED: "No esperado",
}
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
LINE_HEADERS = [
    "Código Producto Largo",
    "BARCODE",
    "Producto",
    "UND",
    "Talla",
    "Color",
    "Según reporte",
    "Vendido en conteo",
    "Esperado",
    "Contado",
    "Diferencia",
    "Estado",
]


def _write_table(ws: Worksheet, headers: list[str], rows: list[list[object]]) -> None:
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[i - 1] or "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 48)


def _line_row(line: ComparisonLine) -> list[object]:
    return [
        line.reference,
        line.ean,
        line.name,
        line.business_unit,
        line.size,
        line.color_name,
        line.reported,
        line.sold,
        line.expected,
        line.counted,
        line.difference,
        STATUS_LABEL[line.status],
    ]


def build_workbook(
    session: CountSession,
    snapshot: StockSnapshot | None,
    summary: ComparisonSummary,
    lines: list[ComparisonLine],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Resumen"
    fmt = "%Y-%m-%d %H:%M"
    info: list[tuple[str, object]] = [
        ("Tienda", f"{session.store.code} - {session.store.name}"),
        ("Conteo", session.name),
        ("Iniciado", session.created_at.strftime(fmt)),
        ("Existencias comparadas", snapshot.effective_at.strftime(fmt) if snapshot else "—"),
        ("Archivo de existencias", snapshot.source_filename if snapshot else "—"),
        ("Generado", datetime.now().strftime(fmt)),
        ("", ""),
        ("Unidades esperadas", summary.expected_units),
        ("Unidades vendidas durante el conteo", summary.sold_units),
        ("Unidades contadas (en tienda)", summary.counted_units),
        ("Unidades cuadradas", summary.matched_units),
        ("Exactitud", f"{summary.accuracy:.1%}"),
        ("Referencias esperadas", summary.expected_lines),
        ("Referencias sin código de barras", summary.lines_without_ean),
    ]
    for status in LineStatus:
        bucket = summary.buckets[status]
        info.append(
            (f"{STATUS_LABEL[status]} (referencias / unidades)", f"{bucket.lines} / {bucket.units}")
        )
    for label, value in info:
        ws.append([label, value])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 48

    _write_table(
        wb.create_sheet("Diferencias"),
        LINE_HEADERS,
        [_line_row(line) for line in lines if line.status != LineStatus.OK],
    )
    _write_table(wb.create_sheet("Todo"), LINE_HEADERS, [_line_row(line) for line in lines])
    # Hoja simple de totales (BARCODE = EAN-13 de la etiqueta).
    _write_table(
        wb.create_sheet("Conteo"),
        ["CÓDIGO PRODUCTO LARGO", "BARCODE", "NOMBRE PRODUCTO", "CONTEO FINAL"],
        [[line.reference, line.ean, line.name, line.counted] for line in lines if line.counted],
    )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
